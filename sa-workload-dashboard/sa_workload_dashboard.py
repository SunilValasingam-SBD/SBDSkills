"""Build an interactive dashboard comparing, per Solution Architect (SA):

  PANEL 1 (left):  count of Worklist items by classification
                   (Project / Idea / TechPass / SA NPD / NPD), restricted to
                   items that are WIP, or Complete items closed within a
                   trailing lookback window (default 28 days).
  PANEL 2 (right): total hours spent by classification
                   (Project / Idea / Time Off / Training & Development /
                   Admin-Other), derived from a timesheet export.
  PANEL 3 (bottom): actual + imputed time allocation by classification
                   (Project / Idea / TechPass / Admin-Other / Training &
                   Development), with each segment labeled with hours AND
                   % of that SA's total. Time Off is excluded entirely.

Worklist classification is derived from the `Work Ref` column prefix:
  PRJ -> Project | ID -> Idea | TechPASS -> TechPass | SA -> SA NPD | NPD -> NPD

Timesheet classification: rows with `Project ID` == the configured time-off
project id (set with `--timeoff-project-id`) are classified by the `Timesheet Line Class`
column (Time Off / Training & Development / everything else -> Admin/Other).
All other rows are classified Project/Idea by their `Project ID` prefix.

Panel 3 exists because Ideas and TechPasses have no dedicated Project ID in
some timesheet systems - that time may be logged under an administrative
project alongside real admin work and time off, so it's invisible in Panel 2.
Panel 3 imputes it instead, based on the Worklist:
  - Every ID-prefixed item accrues `--idea-hours-per-week` hours (default 1)
    for every week (from the timesheet's own set of "Week Start" dates) in
    which it was active.
  - Every TechPASS-prefixed item accrues `--techpass-hours-per-week` hours
    (default 2) the same way.
  - An item is "active" in a given week if its `Created` date (the
    assignment-date proxy) is on/before that week's start AND - for items
    that have reached a terminal status (Complete/Cancelled) - its
    `Modified` date (completion/cancellation-date proxy) is on/after that
    week's start. Still-open items (not yet Complete/Cancelled) are treated
    as active through the end of the available week range.
  - The imputed Idea/TechPass hours are then subtracted from that SA's
    Admin/Other total (floored at 0), since that time is already counted
    there today and would otherwise be double counted.
  - Time Off hours are excluded from Panel 3 entirely (not "work" time).

All three bar charts are horizontal, stacked, color coded with a distinct
high-contrast palette, and show the value directly on each bar segment with
automatic white/black label color (plus an outline) for readability against
any bar color. Panel 3 additionally shows each segment's % of that SA's
total allocated time.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Missing dependency 'openpyxl'. Install with: py -m pip install -r requirements.txt", file=sys.stderr)
    raise

WORKREF_PREFIX_MAP = {
    "PRJ": "Project",
    "ID": "Idea",
    "TECHPASS": "TechPass",
    "SA": "SA NPD",
    "NPD": "NPD",
}

DEFAULT_EXCLUDED_SAS = set()

DEFAULT_LOOKBACK_DAYS = 28
DEFAULT_TIMEOFF_PROJECT_ID = ""
DEFAULT_IDEA_HOURS_PER_WEEK = 1.0
DEFAULT_TECHPASS_HOURS_PER_WEEK = 2.0
TERMINAL_STATUSES = {"Complete", "Cancelled"}

# Distinct, high-contrast qualitative palette so adjacent stacked segments
# never look similar; the muted amber avoids a harsh pure-orange.
PALETTE = [
    "#1f77b4", "#d62728", "#2ca02c", "#9467bd", "#e8974a",
    "#17becf", "#8c564b", "#e377c2", "#bcbd22", "#7f7f7f",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--worklist", type=Path, help="Path to the Worklist CSV export")
    parser.add_argument("--timesheet", type=Path, help="Path to the Timesheet Excel export")
    parser.add_argument(
        "--exclude-sa",
        action="append",
        default=None,
        help="Solution Architect to exclude (exact 'Last, First' match; repeatable). "
             "Blank/omitted = include all architects.",
    )
    parser.add_argument(
        "--lookback-days",
        type=int,
        default=DEFAULT_LOOKBACK_DAYS,
        help=f"Worklist items: include WIP always, plus Complete items closed within this many days (default {DEFAULT_LOOKBACK_DAYS}).",
    )
    parser.add_argument(
        "--timeoff-project-id",
        default=DEFAULT_TIMEOFF_PROJECT_ID,
        help="Timesheet Project ID representing time off/training/admin rows (no default; set for your organization).",
    )
    parser.add_argument(
        "--idea-hours-per-week",
        type=float,
        default=DEFAULT_IDEA_HOURS_PER_WEEK,
        help=f"Imputed hours/week for each active Idea in Panel 3 (default {DEFAULT_IDEA_HOURS_PER_WEEK}).",
    )
    parser.add_argument(
        "--techpass-hours-per-week",
        type=float,
        default=DEFAULT_TECHPASS_HOURS_PER_WEEK,
        help=f"Imputed hours/week for each active TechPass in Panel 3 (default {DEFAULT_TECHPASS_HOURS_PER_WEEK}).",
    )
    parser.add_argument(
        "--no-net-admin",
        action="store_true",
        help="Do not subtract imputed Idea/TechPass hours from the Admin/Other total in Panel 3 (default: subtract).",
    )
    parser.add_argument("--output-dir", type=Path, default=None,
                         help="Output directory (blank = shared SA\\Outputs\\sa-workload-dashboard folder)")
    args = parser.parse_args()

    if args.worklist is None:
        raw = input("Path to the Worklist CSV export: ").strip().strip('"')
        args.worklist = Path(raw)
    if args.timesheet is None:
        raw = input("Path to the Timesheet Excel export: ").strip().strip('"')
        args.timesheet = Path(raw)

    if not args.worklist.exists():
        parser.error(f"Worklist file not found: {args.worklist}")
    if not args.timesheet.exists():
        parser.error(f"Timesheet file not found: {args.timesheet}")

    args.exclude_sa = set(args.exclude_sa) if args.exclude_sa else set(DEFAULT_EXCLUDED_SAS)
    if args.output_dir is None:
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "_shared"))
        from sa_paths import output_dir_for
        args.output_dir = output_dir_for("sa-workload-dashboard")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    return args


def _parse_date(value):
    """Best-effort parse of the mixed date formats found in Worklist exports."""
    if not value:
        return None
    value = str(value).strip()
    if not value:
        return None
    for fmt in ("%m/%d/%Y %I:%M %p", "%m/%d/%Y", "%d-%b-%y", "%d-%b-%Y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def is_recently_active(row, cutoff_date):
    """WIP items are always included; Complete items only if closed within the lookback window."""
    status = (row.get("Status") or "").strip()
    if status == "WIP":
        return True
    if status != "Complete":
        return False
    closure_date = _parse_date(row.get("Work Ref: Finish Date")) or _parse_date(row.get("Modified"))
    if closure_date is None:
        return False
    return closure_date >= cutoff_date


def classify_timeclass(value):
    v = (value or "").strip().lower()
    if "time off" in v:
        return "Time Off"
    if "training" in v:
        return "Training & Development"
    return "Admin / Other"


def load_worklist(path, excluded_sas, cutoff_date):
    """Return {sa: {classification: count}}"""
    counts = defaultdict(lambda: defaultdict(int))
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            sa = (row.get("Work Ref: Solution Architect") or "").strip()
            work_ref = (row.get("Work Ref") or "").strip()
            if not sa or not work_ref or sa in excluded_sas:
                continue
            if not is_recently_active(row, cutoff_date):
                continue
            m = re.match(r"^[A-Za-z]+", work_ref)
            prefix = m.group(0).upper() if m else ""
            classification = WORKREF_PREFIX_MAP.get(prefix, prefix or "Unclassified")
            counts[sa][classification] += 1
    return counts


def load_timesheet(path, excluded_sas, timeoff_project_id):
    """Return ({sa: {classification: hours}}, sorted list of distinct Week Start dates)."""
    hours = defaultdict(lambda: defaultdict(float))
    weeks_seen = set()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    header_row_idx = None
    header = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        if row and row[0] == "Segment Name":
            header_row_idx = i
            header = list(row)
            break
    if header is None:
        raise ValueError("Could not locate header row (expected a 'Segment Name' column) in timesheet export")

    col = {name: idx for idx, name in enumerate(header)}
    required = ["Resource Name", "Project ID", "Timesheet Line Class", "Hours Reported"]
    missing = [c for c in required if c not in col]
    if missing:
        raise ValueError(f"Timesheet export is missing expected column(s): {missing}")

    resource_col = col["Resource Name"]
    project_id_col = col["Project ID"]
    timeclass_col = col["Timesheet Line Class"]
    hours_col = col["Hours Reported"]
    week_start_col = col.get("Week Start")

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None or row[resource_col] is None:
            continue
        sa = (row[resource_col] or "").strip()
        project_id = (row[project_id_col] or "").strip()
        reported_hours = row[hours_col] or 0
        if not sa or not project_id or sa in excluded_sas:
            continue

        if project_id == timeoff_project_id:
            classification = classify_timeclass(row[timeclass_col])
        else:
            m = re.match(r"^[A-Za-z]+", project_id)
            prefix = m.group(0).upper() if m else ""
            if prefix == "PRJ":
                classification = "Project"
            elif prefix == "ID":
                classification = "Idea"
            else:
                classification = prefix or "Unclassified"

        hours[sa][classification] += float(reported_hours)

        if week_start_col is not None:
            week_val = row[week_start_col]
            week_date = week_val if isinstance(week_val, datetime) else _parse_date(week_val)
            if week_date is not None:
                weeks_seen.add(week_date.date())

    return hours, sorted(weeks_seen)


def load_impute_items(path, excluded_sas):
    """Return a list of {sa, classification ('Idea'/'TechPass'), created, end_date} for
    every Idea/TechPass Worklist row with a usable `Created` date. `end_date` is None for
    items that have not reached a terminal status (Complete/Cancelled) - these are treated
    as active through the end of the available week range."""
    items = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            sa = (row.get("Work Ref: Solution Architect") or "").strip()
            work_ref = (row.get("Work Ref") or "").strip()
            if not sa or not work_ref or sa in excluded_sas:
                continue
            m = re.match(r"^[A-Za-z]+", work_ref)
            prefix = m.group(0).upper() if m else ""
            classification = WORKREF_PREFIX_MAP.get(prefix)
            if classification not in ("Idea", "TechPass"):
                continue
            created = _parse_date(row.get("Created"))
            if created is None:
                continue
            status = (row.get("Status") or "").strip()
            end_date = _parse_date(row.get("Modified")) if status in TERMINAL_STATUSES else None
            items.append({
                "sa": sa,
                "classification": classification,
                "created": created.date(),
                "end_date": end_date.date() if end_date else None,
            })
    return items


def compute_imputed_hours(items, weeks, idea_hours_per_week, techpass_hours_per_week):
    """Return {sa: {'Idea'/'TechPass': total imputed hours}} by accruing hours/week for
    every week (from `weeks`, a sorted list of date objects) that each item was active in."""
    hours = defaultdict(lambda: defaultdict(float))
    if not weeks:
        return hours
    per_week = {"Idea": idea_hours_per_week, "TechPass": techpass_hours_per_week}
    for item in items:
        rate = per_week[item["classification"]]
        for week in weeks:
            if week < item["created"]:
                continue
            if item["end_date"] is not None and week > item["end_date"]:
                continue
            hours[item["sa"]][item["classification"]] += rate
    return hours


def build_dataset(counts_by_sa, categories_order, sas):
    datasets = []
    for cat in categories_order:
        datasets.append({
            "label": cat,
            "data": [round(counts_by_sa.get(sa, {}).get(cat, 0), 2) for sa in sas],
        })
    return datasets


def build_dataset_with_pct(counts_by_sa, categories_order, sas):
    """Like build_dataset, but each dataset also carries a parallel `percentages` array -
    each value's share of that SA's total across categories_order - for on-bar % labels."""
    totals = {sa: sum(counts_by_sa.get(sa, {}).get(cat, 0.0) for cat in categories_order) for sa in sas}
    datasets = []
    for cat in categories_order:
        data = [round(counts_by_sa.get(sa, {}).get(cat, 0.0), 2) for sa in sas]
        pct = [round((counts_by_sa.get(sa, {}).get(cat, 0.0) / totals[sa] * 100), 1) if totals[sa] else 0.0 for sa in sas]
        datasets.append({"label": cat, "data": data, "percentages": pct})
    return datasets


def colorize(datasets):
    out = []
    for i, ds in enumerate(datasets):
        ds = dict(ds)
        ds["backgroundColor"] = PALETTE[i % len(PALETTE)]
        out.append(ds)
    return out


def build_html(sas, left_datasets, right_datasets, third_datasets, impute_note):
    left_datasets = colorize(left_datasets)
    right_datasets = colorize(right_datasets)
    third_datasets = colorize(third_datasets)
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>SA Workload Dashboard</title>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4"></script>
<script src="https://cdn.jsdelivr.net/npm/chartjs-plugin-datalabels@2"></script>
<style>
  body {{ font-family: Segoe UI, Arial, sans-serif; margin: 20px; background: #f7f8fa; }}
  h1 {{ text-align: center; }}
  .container {{ display: flex; gap: 20px; flex-wrap: wrap; }}
  .panel {{ flex: 1; min-width: 480px; background: white; border-radius: 8px; padding: 16px;
            box-shadow: 0 1px 4px rgba(0,0,0,0.15); }}
  .panel.full {{ flex-basis: 100%; }}
  .panel h2 {{ text-align: center; font-size: 18px; margin-bottom: 4px; }}
  .panel .note {{ text-align: center; font-size: 12px; color: #666; margin-bottom: 10px; }}
  canvas {{ max-height: 650px; }}
</style>
</head>
<body>
<h1>Solution Architect Workload: Worklist Classification vs. Timesheet Hours</h1>
<div class="container">
  <div class="panel">
    <h2>Worklist Item Count by Classification</h2>
    <canvas id="leftChart"></canvas>
  </div>
  <div class="panel">
    <h2>Hours Spent by Classification</h2>
    <canvas id="rightChart"></canvas>
  </div>
  <div class="panel full">
    <h2>Time Allocation by Classification (Actual + Imputed, with % of total)</h2>
    <div class="note">{impute_note}</div>
    <canvas id="thirdChart"></canvas>
  </div>
</div>
<script>
const sas = {json.dumps(sas)};
const leftDatasets = {json.dumps(left_datasets)};
const rightDatasets = {json.dumps(right_datasets)};
const thirdDatasets = {json.dumps(third_datasets)};

function contrastColor(hex, invert) {{
  hex = hex.replace('#', '');
  const r = parseInt(hex.substring(0, 2), 16);
  const g = parseInt(hex.substring(2, 4), 16);
  const b = parseInt(hex.substring(4, 6), 16);
  const luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255;
  const light = luminance > 0.55;
  if (invert) return light ? '#ffffff' : '#000000';
  return light ? '#000000' : '#ffffff';
}}

function makeChart(canvasId, labels, datasets, yLabel, withPct) {{
  new Chart(document.getElementById(canvasId), {{
    type: 'bar',
    data: {{ labels: labels, datasets: datasets }},
    plugins: [ChartDataLabels],
    options: {{
      indexAxis: 'y',
      responsive: true,
      maintainAspectRatio: false,
      plugins: {{
        legend: {{ position: 'bottom' }},
        tooltip: {{
          mode: 'index',
          intersect: false,
          callbacks: withPct ? {{
            label: function(context) {{
              const pct = context.dataset.percentages ? context.dataset.percentages[context.dataIndex] : null;
              const base = context.dataset.label + ': ' + context.formattedValue;
              return pct !== null ? base + ' (' + pct.toFixed(1) + '%)' : base;
            }}
          }} : undefined
        }},
        datalabels: {{
          color: function(context) {{ return contrastColor(context.dataset.backgroundColor, false); }},
          textStrokeColor: function(context) {{ return contrastColor(context.dataset.backgroundColor, true); }},
          textStrokeWidth: 2,
          font: {{ size: 11, weight: 'bold' }},
          formatter: function(value, context) {{
            if (!(value > 0)) return '';
            if (withPct && context.dataset.percentages) {{
              const pct = context.dataset.percentages[context.dataIndex];
              return value + 'h (' + pct.toFixed(0) + '%)';
            }}
            return value;
          }}
        }}
      }},
      scales: {{
        x: {{ stacked: true, title: {{ display: true, text: yLabel }} }},
        y: {{ stacked: true }}
      }}
    }}
  }});
}}
makeChart('leftChart', sas, leftDatasets, 'Number of Items', false);
makeChart('rightChart', sas, rightDatasets, 'Hours', false);
makeChart('thirdChart', sas, thirdDatasets, 'Hours', true);
</script>
</body>
</html>
"""


def main():
    args = parse_args()
    cutoff_date = datetime.now() - timedelta(days=args.lookback_days)

    worklist_counts = load_worklist(args.worklist, args.exclude_sa, cutoff_date)
    timesheet_hours, timesheet_weeks = load_timesheet(args.timesheet, args.exclude_sa, args.timeoff_project_id)
    impute_items = load_impute_items(args.worklist, args.exclude_sa)
    impute_hours = compute_imputed_hours(
        impute_items, timesheet_weeks, args.idea_hours_per_week, args.techpass_hours_per_week
    )

    worklist_categories = ["Project", "Idea", "TechPass", "SA NPD", "NPD"]
    extra = sorted({c for d in worklist_counts.values() for c in d if c not in worklist_categories})
    worklist_categories += extra

    timesheet_categories = ["Project", "Idea", "Time Off", "Training & Development", "Admin / Other"]
    extra2 = sorted({c for d in timesheet_hours.values() for c in d if c not in timesheet_categories})
    timesheet_categories += extra2

    all_sas = sorted(set(worklist_counts.keys()) | set(timesheet_hours.keys()) | set(impute_hours.keys()))

    left_datasets = build_dataset(worklist_counts, worklist_categories, all_sas)
    right_datasets = build_dataset(timesheet_hours, timesheet_categories, all_sas)

    # Panel 3: actual Project hours + imputed Idea/TechPass hours, Admin/Other netted
    # of the imputed hours it already implicitly contains, Time Off excluded entirely.
    third_categories = ["Project", "Idea", "TechPass", "Admin / Other", "Training & Development"]
    third_counts = defaultdict(lambda: defaultdict(float))
    for sa in all_sas:
        project_hrs = timesheet_hours.get(sa, {}).get("Project", 0.0)
        idea_hrs = impute_hours.get(sa, {}).get("Idea", 0.0)
        techpass_hrs = impute_hours.get(sa, {}).get("TechPass", 0.0)
        admin_hrs = timesheet_hours.get(sa, {}).get("Admin / Other", 0.0)
        if not args.no_net_admin:
            admin_hrs = max(0.0, admin_hrs - idea_hrs - techpass_hrs)
        training_hrs = timesheet_hours.get(sa, {}).get("Training & Development", 0.0)
        third_counts[sa]["Project"] = project_hrs
        third_counts[sa]["Idea"] = idea_hrs
        third_counts[sa]["TechPass"] = techpass_hrs
        third_counts[sa]["Admin / Other"] = admin_hrs
        third_counts[sa]["Training & Development"] = training_hrs

    third_datasets = build_dataset_with_pct(third_counts, third_categories, all_sas)

    week_range = f"{timesheet_weeks[0]} to {timesheet_weeks[-1]}" if timesheet_weeks else "no weeks found in timesheet"
    net_desc = "netted out of Admin/Other" if not args.no_net_admin else "NOT netted out of Admin/Other"
    impute_note = (
        f"Imputed: {args.idea_hours_per_week}h/week per active Idea, {args.techpass_hours_per_week}h/week per active "
        f"TechPass, accrued over {len(timesheet_weeks)} week(s) in the timesheet range ({week_range}); {net_desc}. "
        f"Time Off excluded."
    )

    output_html = args.output_dir / "sa_dashboard.html"
    html = build_html(all_sas, left_datasets, right_datasets, third_datasets, impute_note)
    output_html.write_text(html, encoding="utf-8")

    print(f"Wrote {output_html}")
    print(f"Lookback cutoff date: {cutoff_date.date()} ({args.lookback_days} days)")
    print(f"Excluded SAs: {sorted(args.exclude_sa)}")
    print(f"SAs included ({len(all_sas)}): {all_sas}")
    print(f"Timesheet week range for Panel 3 imputation: {week_range} ({len(timesheet_weeks)} week(s))")


if __name__ == "__main__":
    main()
