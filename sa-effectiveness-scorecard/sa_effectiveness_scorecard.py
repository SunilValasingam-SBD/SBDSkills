"""Build an interactive "SA Effectiveness Scorecard" per Solution Architect (SA),
combining a Worklist CSV export with a Timesheet Excel export to show how much of
each SA's time is tangible delivery work vs. internal/NPD work vs. admin overhead.

WHY THIS EXISTS
----------------
sa-workload-dashboard already shows Worklist counts next to Timesheet hours, but it
treats every classification as its own bucket with no overall "how productive is
this SA" signal. This skill goes one step further: it buckets every classification
into one of THREE TIERS and ranks SAs by the % of their time that is genuinely
tangible, delivery-facing work - their "brand value" - rather than just counting
tickets or admin hours.

THE THREE TIERS
----------------
  TANGIBLE / DELIVERY  - Project (actual timesheet hours) + Idea (imputed) +
                          TechPass (imputed). Real deliverables with business value.
  INTERNAL / NPD        - SA NPD (imputed) + NPD (imputed). Legitimate SA work, but
                          not SARB/SLA-tracked delivery - kept as its own tier so it
                          is never silently folded into either "productive" or
                          "overhead" and mis-read either way.
  ADMIN / OVERHEAD       - Admin/Other (netted of all imputed hours below) +
                          Training & Development. Time Off is excluded entirely -
                          it is not work time at all.

Effectiveness Score (the headline ranking metric) = Tangible hours / (Tangible +
Internal/NPD + Admin) hours, as a %. SAs are ranked highest-to-lowest.

IMPUTATION (why it's needed)
------------------------------
Idea, TechPass, SA NPD, and NPD Worklist items have NO dedicated Project ID in the
timesheet system - that time may be logged under an administrative project
alongside real admin work and time off, so it is invisible unless imputed. This
script imputes it from the Worklist, exactly like sa-workload-dashboard
does for Idea/TechPass, but extended to also cover SA NPD and NPD:
  - Every Idea-prefixed item accrues `--idea-hours-per-week` hours/week (default 1).
  - Every TechPass-prefixed item accrues `--techpass-hours-per-week` hours/week (default 2).
  - Every SA NPD-prefixed item accrues `--sa-npd-hours-per-week` hours/week (default 1).
  - Every NPD-prefixed item accrues `--npd-hours-per-week` hours/week (default 1).
  - "Weeks" = the distinct Week Start values present in the Timesheet export.
  - An item is "active" in a week if its Created date is on/before that week's start
    AND - if it has reached a terminal status (Complete/Cancelled) - its Modified
    date is on/after that week's start. Non-terminal items are active through the
    end of the available week range.
  - All four imputed totals are subtracted from that SA's Admin/Other hours
    (floored at 0) to avoid double-counting, since that time is already implicitly
    sitting in Admin/Other today.

All imputed rates are fully configurable (see --help) - adjust them if your org's
typical effort-per-item differs.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Missing dependency 'openpyxl'. Install with: py -m pip install -r requirements.txt", file=sys.stderr)
    raise

WORKREF_PREFIX_MAP = {
    "PRJ": "Project",
    "ID": "Idea",
    "TECHPASS": "TechPass",
    "SA": "SA NPD",
    "NPD": "NPD",
}

IMPUTED_CLASSIFICATIONS = ("Idea", "TechPass", "SA NPD", "NPD")

TIER_MAP = {
    "Project": "Tangible / Delivery",
    "Idea": "Tangible / Delivery",
    "TechPass": "Tangible / Delivery",
    "SA NPD": "Internal / NPD",
    "NPD": "Internal / NPD",
    "Admin / Other": "Admin / Overhead",
    "Training & Development": "Admin / Overhead",
}
TIER_ORDER = ["Tangible / Delivery", "Internal / NPD", "Admin / Overhead"]
TIER_COLORS = {
    "Tangible / Delivery": "#2ca02c",   # green
    "Internal / NPD": "#1f77b4",        # blue
    "Admin / Overhead": "#e8974a",      # amber
}
TIER_TAGS = {
    "Tangible / Delivery": "\U0001F7E2 Tangible / Delivery",
    "Internal / NPD": "\U0001F535 Internal / NPD",
    "Admin / Overhead": "\U0001F7E0 Admin / Overhead",
}

DEFAULT_EXCLUDED_SAS = set()

DEFAULT_TIMEOFF_PROJECT_ID = ""
DEFAULT_IDEA_HOURS_PER_WEEK = 1.0
DEFAULT_TECHPASS_HOURS_PER_WEEK = 2.0
DEFAULT_SA_NPD_HOURS_PER_WEEK = 1.0
DEFAULT_NPD_HOURS_PER_WEEK = 1.0
TERMINAL_STATUSES = {"Complete", "Cancelled"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--worklist", type=Path, help="Path to the Worklist CSV export")
    parser.add_argument("--timesheet", type=Path, help="Path to the Timesheet Excel export")
    parser.add_argument(
        "--exclude-sa",
        action="append",
        default=None,
        help="Solution Architect to exclude (exact 'Last, First' match; repeatable). "
             "Blank/omitted = include all architects.",
    )
    parser.add_argument(
        "--timeoff-project-id",
        default=DEFAULT_TIMEOFF_PROJECT_ID,
        help="Timesheet Project ID representing time off/training/admin rows (no default; set for your organization).",
    )
    parser.add_argument("--idea-hours-per-week", type=float, default=DEFAULT_IDEA_HOURS_PER_WEEK,
                         help=f"Imputed hours/week per active Idea (default {DEFAULT_IDEA_HOURS_PER_WEEK}).")
    parser.add_argument("--techpass-hours-per-week", type=float, default=DEFAULT_TECHPASS_HOURS_PER_WEEK,
                         help=f"Imputed hours/week per active TechPass (default {DEFAULT_TECHPASS_HOURS_PER_WEEK}).")
    parser.add_argument("--sa-npd-hours-per-week", type=float, default=DEFAULT_SA_NPD_HOURS_PER_WEEK,
                         help=f"Imputed hours/week per active SA NPD item (default {DEFAULT_SA_NPD_HOURS_PER_WEEK}).")
    parser.add_argument("--npd-hours-per-week", type=float, default=DEFAULT_NPD_HOURS_PER_WEEK,
                         help=f"Imputed hours/week per active NPD item (default {DEFAULT_NPD_HOURS_PER_WEEK}).")
    parser.add_argument(
        "--no-net-admin",
        action="store_true",
        help="Do not subtract imputed Idea/TechPass/SA NPD/NPD hours from Admin/Other (default: subtract).",
    )
    parser.add_argument("--output-dir", type=Path, default=None,
                         help="Output directory (blank = shared SA\\Outputs\\sa-effectiveness-scorecard folder)")
    args = parser.parse_args()

    if args.worklist is None:
        raw = input("Path to the Worklist CSV export: ").strip().strip('"')
        args.worklist = Path(raw)
    if args.timesheet is None:
        raw = input("Path to the Timesheet Excel export: ").strip().strip('"')
        args.timesheet = Path(raw)

    if not args.worklist.exists():
        parser.error(f"Worklist file not found: {args.worklist}")
    if not args.timesheet.exists():
        parser.error(f"Timesheet file not found: {args.timesheet}")

    args.exclude_sa = set(args.exclude_sa) if args.exclude_sa else set(DEFAULT_EXCLUDED_SAS)
    if args.output_dir is None:
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "_shared"))
        from sa_paths import output_dir_for
        args.output_dir = output_dir_for("sa-effectiveness-scorecard")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    return args


def _parse_date(value):
    if not value:
        return None
    value = str(value).strip()
    if not value:
        return None
    for fmt in ("%m/%d/%Y %I:%M %p", "%m/%d/%Y", "%d-%b-%y", "%d-%b-%Y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def classify_timeclass(value):
    v = (value or "").strip().lower()
    if "time off" in v:
        return "Time Off"
    if "training" in v:
        return "Training & Development"
    return "Admin / Other"


def load_worklist_items(path, excluded_sas):
    """Return every classified Worklist row, keyed by SA, with its classification,
    Work Ref / name, Created, Status, Modified - used both for imputation and for
    the per-SA drill-down tables (all items, not just a recency window - this skill
    is about overall time allocation, not "currently active" ticket counts)."""
    items = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            sa = (row.get("Work Ref: Solution Architect") or "").strip()
            work_ref = (row.get("Work Ref") or "").strip()
            if not sa or not work_ref or sa in excluded_sas:
                continue
            m = re.match(r"^[A-Za-z]+", work_ref)
            prefix = m.group(0).upper() if m else ""
            classification = WORKREF_PREFIX_MAP.get(prefix, prefix or "Unclassified")
            created = _parse_date(row.get("Created"))
            status = (row.get("Status") or "").strip()
            modified = _parse_date(row.get("Modified"))
            items.append({
                "sa": sa,
                "work_ref": work_ref,
                "classification": classification,
                "status": status,
                "created": created,
                "modified": modified,
            })
    return items


def load_timesheet(path, excluded_sas, timeoff_project_id):
    hours = defaultdict(lambda: defaultdict(float))
    weeks_seen = set()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    header_row_idx = None
    header = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        if row and row[0] == "Segment Name":
            header_row_idx = i
            header = list(row)
            break
    if header is None:
        raise ValueError("Could not locate header row (expected a 'Segment Name' column) in timesheet export")

    col = {name: idx for idx, name in enumerate(header)}
    required = ["Resource Name", "Project ID", "Timesheet Line Class", "Hours Reported"]
    missing = [c for c in required if c not in col]
    if missing:
        raise ValueError(f"Timesheet export is missing expected column(s): {missing}")

    resource_col = col["Resource Name"]
    project_id_col = col["Project ID"]
    timeclass_col = col["Timesheet Line Class"]
    hours_col = col["Hours Reported"]
    week_start_col = col.get("Week Start")

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None or row[resource_col] is None:
            continue
        sa = (row[resource_col] or "").strip()
        project_id = (row[project_id_col] or "").strip()
        reported_hours = row[hours_col] or 0
        if not sa or not project_id or sa in excluded_sas:
            continue

        if project_id == timeoff_project_id:
            classification = classify_timeclass(row[timeclass_col])
        else:
            m = re.match(r"^[A-Za-z]+", project_id)
            prefix = m.group(0).upper() if m else ""
            classification = "Project" if prefix == "PRJ" else ("Idea" if prefix == "ID" else (prefix or "Unclassified"))

        hours[sa][classification] += float(reported_hours)

        if week_start_col is not None:
            week_val = row[week_start_col]
            week_date = week_val if isinstance(week_val, datetime) else _parse_date(week_val)
            if week_date is not None:
                weeks_seen.add(week_date.date())

    return hours, sorted(weeks_seen)


def compute_imputed_hours(items, weeks, rates):
    """Return {sa: {classification: total imputed hours}} for Idea/TechPass/SA NPD/NPD."""
    hours = defaultdict(lambda: defaultdict(float))
    if not weeks:
        return hours
    for item in items:
        classification = item["classification"]
        if classification not in IMPUTED_CLASSIFICATIONS:
            continue
        created = item["created"]
        if created is None:
            continue
        created = created.date()
        end_date = item["modified"].date() if (item["status"] in TERMINAL_STATUSES and item["modified"]) else None
        rate = rates[classification]
        for week in weeks:
            if week < created:
                continue
            if end_date is not None and week > end_date:
                continue
            hours[item["sa"]][classification] += rate
    return hours


def build_html(sas, tier_datasets, scorecard_rows, drilldown, impute_note):
    tier_json = json.dumps(tier_datasets)
    scorecard_json = json.dumps(scorecard_rows)
    drilldown_json = json.dumps(drilldown)
    sas_json = json.dumps(sas)
    tier_colors_json = json.dumps(TIER_COLORS)
    tier_tags_json = json.dumps(TIER_TAGS)
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>SA Effectiveness Scorecard</title>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4"></script>
<script src="https://cdn.jsdelivr.net/npm/chartjs-plugin-datalabels@2"></script>
<style>
  body {{ font-family: Segoe UI, Arial, sans-serif; margin: 20px; background: #f7f8fa; }}
  h1 {{ text-align: center; }}
  .legend-tags {{ text-align: center; margin: 6px 0 18px; }}
  .tag {{ display: inline-block; padding: 4px 12px; border-radius: 14px; margin: 0 6px; font-size: 13px;
          font-weight: 600; color: #222; background: #eee; border: 1px solid #ccc; }}
  .panel {{ background: white; border-radius: 8px; padding: 16px; margin-bottom: 20px;
            box-shadow: 0 1px 4px rgba(0,0,0,0.15); }}
  .panel h2 {{ text-align: center; font-size: 18px; margin-bottom: 4px; }}
  .panel .note {{ text-align: center; font-size: 12px; color: #666; margin-bottom: 10px; }}
  canvas {{ max-height: 750px; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
  th, td {{ border: 1px solid #ddd; padding: 6px 8px; text-align: right; }}
  th {{ background: #f0f2f5; position: sticky; top: 0; }}
  td:first-child, th:first-child {{ text-align: left; }}
  tr:nth-child(even) {{ background: #fafbfc; }}
  tr.clickable-row {{ cursor: pointer; }}
  tr.clickable-row:hover {{ background: #eef5ff; }}
  .rank-1 {{ background: #fff9db !important; }}
  #drilldownWrap {{ display: none; }}
  .back-link {{ cursor: pointer; color: #1f77b4; font-size: 13px; margin-bottom: 8px; display: inline-block; }}
  .score-pill {{ font-weight: bold; padding: 2px 8px; border-radius: 10px; color: white; }}
</style>
</head>
<body>
<h1>Solution Architect Effectiveness Scorecard</h1>
<div class="legend-tags" id="legendTags"></div>

<div class="panel">
  <h2>Effectiveness Leaderboard (ranked by Tangible/Delivery %)</h2>
  <div class="note">{impute_note}</div>
  <canvas id="leaderboardChart"></canvas>
</div>

<div class="panel">
  <h2>Scorecard (click a row to see that SA's underlying Worklist items)</h2>
  <table id="scorecardTable">
    <thead>
      <tr>
        <th>Rank</th><th>Solution Architect</th>
        <th>Total Hours</th>
        <th>Tangible / Delivery h</th><th>Tangible %</th>
        <th>Internal / NPD h</th><th>Internal %</th>
        <th>Admin / Overhead h</th><th>Admin %</th>
        <th>Effectiveness Score</th>
      </tr>
    </thead>
    <tbody id="scorecardBody"></tbody>
  </table>
</div>

<div class="panel" id="drilldownWrap">
  <span class="back-link" onclick="document.getElementById('drilldownWrap').style.display='none';">&#8592; Hide</span>
  <h2 id="drilldownTitle"></h2>
  <table id="drilldownTable">
    <thead><tr><th>Work Ref</th><th>Classification</th><th>Tier</th><th>Status</th><th>Created</th></tr></thead>
    <tbody id="drilldownBody"></tbody>
  </table>
</div>

<script>
const sas = {sas_json};
const tierDatasets = {tier_json};
const scorecardRows = {scorecard_json};
const drilldown = {drilldown_json};
const TIER_COLORS = {tier_colors_json};
const TIER_TAGS = {tier_tags_json};

document.getElementById('legendTags').innerHTML = Object.values(TIER_TAGS).map((t, i) => {{
  const color = Object.values(TIER_COLORS)[i];
  return `<span class="tag" style="border-left: 6px solid ${{color}};">${{t}}</span>`;
}}).join('');

function contrastColor(hex) {{
  hex = hex.replace('#', '');
  const r = parseInt(hex.substring(0, 2), 16);
  const g = parseInt(hex.substring(2, 4), 16);
  const b = parseInt(hex.substring(4, 6), 16);
  const luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255;
  return luminance > 0.55 ? '#000000' : '#ffffff';
}}

const coloredDatasets = tierDatasets.map(ds => ({{
  ...ds,
  backgroundColor: TIER_COLORS[ds.label],
}}));

new Chart(document.getElementById('leaderboardChart'), {{
  type: 'bar',
  data: {{ labels: sas, datasets: coloredDatasets }},
  plugins: [ChartDataLabels],
  options: {{
    indexAxis: 'y',
    responsive: true,
    maintainAspectRatio: false,
    plugins: {{
      legend: {{ position: 'bottom' }},
      tooltip: {{
        mode: 'index',
        intersect: false,
        callbacks: {{
          label: function(context) {{
            const pct = context.dataset.percentages ? context.dataset.percentages[context.dataIndex] : null;
            const base = context.dataset.label + ': ' + context.formattedValue + 'h';
            return pct !== null ? base + ' (' + pct.toFixed(1) + '%)' : base;
          }}
        }}
      }},
      datalabels: {{
        color: function(context) {{ return contrastColor(context.dataset.backgroundColor); }},
        font: {{ size: 11, weight: 'bold' }},
        formatter: function(value, context) {{
          if (!(value > 0)) return '';
          const pct = context.dataset.percentages[context.dataIndex];
          return pct.toFixed(0) + '%';
        }}
      }}
    }},
    scales: {{
      x: {{ stacked: true, title: {{ display: true, text: 'Hours' }} }},
      y: {{ stacked: true }}
    }}
  }}
}});

function renderScorecard() {{
  const body = document.getElementById('scorecardBody');
  body.innerHTML = scorecardRows.map((r, i) => `
    <tr class="clickable-row ${{i === 0 ? 'rank-1' : ''}}" onclick="showDrilldown('${{r.sa.replace(/'/g, "\\\\'")}}')">
      <td>${{i + 1}}</td>
      <td>${{r.sa}}</td>
      <td>${{r.total.toFixed(1)}}</td>
      <td>${{r.tangible.toFixed(1)}}</td><td>${{r.tangible_pct.toFixed(1)}}%</td>
      <td>${{r.internal.toFixed(1)}}</td><td>${{r.internal_pct.toFixed(1)}}%</td>
      <td>${{r.admin.toFixed(1)}}</td><td>${{r.admin_pct.toFixed(1)}}%</td>
      <td><span class="score-pill" style="background:${{r.tangible_pct >= 50 ? '#2ca02c' : (r.tangible_pct >= 25 ? '#e8974a' : '#d62728')}}">${{r.tangible_pct.toFixed(0)}}%</span></td>
    </tr>
  `).join('');
}}
renderScorecard();

function showDrilldown(sa) {{
  const wrap = document.getElementById('drilldownWrap');
  const items = drilldown[sa] || [];
  document.getElementById('drilldownTitle').textContent = sa + ' — ' + items.length + ' Worklist item(s)';
  document.getElementById('drilldownBody').innerHTML = items.map(it => `
    <tr><td>${{it.work_ref}}</td><td>${{it.classification}}</td><td>${{it.tier}}</td><td>${{it.status}}</td><td>${{it.created || ''}}</td></tr>
  `).join('');
  wrap.style.display = 'block';
  wrap.scrollIntoView({{ behavior: 'smooth' }});
}}
</script>
</body>
</html>
"""


def main():
    args = parse_args()

    items = load_worklist_items(args.worklist, args.exclude_sa)
    timesheet_hours, timesheet_weeks = load_timesheet(args.timesheet, args.exclude_sa, args.timeoff_project_id)
    rates = {
        "Idea": args.idea_hours_per_week,
        "TechPass": args.techpass_hours_per_week,
        "SA NPD": args.sa_npd_hours_per_week,
        "NPD": args.npd_hours_per_week,
    }
    impute_hours = compute_imputed_hours(items, timesheet_weeks, rates)

    all_sas = sorted(set(timesheet_hours.keys()) | set(impute_hours.keys()) | {it["sa"] for it in items})

    # Build per-SA, per-classification hours, then roll up into the three tiers.
    tier_hours = defaultdict(lambda: defaultdict(float))
    for sa in all_sas:
        project_hrs = timesheet_hours.get(sa, {}).get("Project", 0.0)
        idea_hrs = impute_hours.get(sa, {}).get("Idea", 0.0)
        techpass_hrs = impute_hours.get(sa, {}).get("TechPass", 0.0)
        sa_npd_hrs = impute_hours.get(sa, {}).get("SA NPD", 0.0)
        npd_hrs = impute_hours.get(sa, {}).get("NPD", 0.0)
        admin_hrs = timesheet_hours.get(sa, {}).get("Admin / Other", 0.0)
        training_hrs = timesheet_hours.get(sa, {}).get("Training & Development", 0.0)

        if not args.no_net_admin:
            admin_hrs = max(0.0, admin_hrs - idea_hrs - techpass_hrs - sa_npd_hrs - npd_hrs)

        tier_hours[sa]["Tangible / Delivery"] = project_hrs + idea_hrs + techpass_hrs
        tier_hours[sa]["Internal / NPD"] = sa_npd_hrs + npd_hrs
        tier_hours[sa]["Admin / Overhead"] = admin_hrs + training_hrs

    # Scorecard rows, ranked by Tangible % (Effectiveness Score) descending.
    scorecard_rows = []
    for sa in all_sas:
        th = tier_hours[sa]
        total = sum(th.values())
        tangible = th["Tangible / Delivery"]
        internal = th["Internal / NPD"]
        admin = th["Admin / Overhead"]
        tangible_pct = (tangible / total * 100) if total else 0.0
        internal_pct = (internal / total * 100) if total else 0.0
        admin_pct = (admin / total * 100) if total else 0.0
        scorecard_rows.append({
            "sa": sa, "total": round(total, 1),
            "tangible": round(tangible, 1), "tangible_pct": round(tangible_pct, 1),
            "internal": round(internal, 1), "internal_pct": round(internal_pct, 1),
            "admin": round(admin, 1), "admin_pct": round(admin_pct, 1),
        })
    scorecard_rows.sort(key=lambda r: r["tangible_pct"], reverse=True)
    ranked_sas = [r["sa"] for r in scorecard_rows]

    # Leaderboard chart datasets follow the same SA order as the scorecard.
    tier_datasets = []
    for tier in TIER_ORDER:
        data = [round(tier_hours[sa][tier], 1) for sa in ranked_sas]
        totals = [sum(tier_hours[sa].values()) for sa in ranked_sas]
        pct = [round((d / t * 100), 1) if t else 0.0 for d, t in zip(data, totals)]
        tier_datasets.append({"label": tier, "data": data, "percentages": pct})

    # Drill-down: every Worklist item per SA, tagged with its tier.
    drilldown = defaultdict(list)
    for it in items:
        tier = TIER_MAP.get(it["classification"])
        if tier is None:
            continue
        drilldown[it["sa"]].append({
            "work_ref": it["work_ref"],
            "classification": it["classification"],
            "tier": tier,
            "status": it["status"],
            "created": it["created"].strftime("%Y-%m-%d") if it["created"] else "",
        })
    for sa in drilldown:
        drilldown[sa].sort(key=lambda r: (r["tier"], r["work_ref"]))

    week_range = f"{timesheet_weeks[0]} to {timesheet_weeks[-1]}" if timesheet_weeks else "no weeks found in timesheet"
    net_desc = "netted out of Admin/Other" if not args.no_net_admin else "NOT netted out of Admin/Other"
    impute_note = (
        f"Imputed: {args.idea_hours_per_week}h/wk/Idea, {args.techpass_hours_per_week}h/wk/TechPass, "
        f"{args.sa_npd_hours_per_week}h/wk/SA NPD, {args.npd_hours_per_week}h/wk/NPD, accrued over "
        f"{len(timesheet_weeks)} week(s) in the timesheet range ({week_range}); {net_desc}. Time Off excluded."
    )

    html = build_html(ranked_sas, tier_datasets, scorecard_rows, dict(drilldown), impute_note)
    (args.output_dir / "sa-effectiveness-scorecard.html").write_text(html, encoding="utf-8")

    summary_csv = args.output_dir / "sa-effectiveness-summary.csv"
    with open(summary_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([
            "Rank", "Solution Architect", "Total Hours",
            "Tangible/Delivery Hours", "Tangible %",
            "Internal/NPD Hours", "Internal %",
            "Admin/Overhead Hours", "Admin %", "Effectiveness Score",
        ])
        for i, r in enumerate(scorecard_rows, start=1):
            writer.writerow([
                i, r["sa"], r["total"], r["tangible"], r["tangible_pct"],
                r["internal"], r["internal_pct"], r["admin"], r["admin_pct"], r["tangible_pct"],
            ])

    print(f"Wrote {args.output_dir / 'sa-effectiveness-scorecard.html'}")
    print(f"Wrote {summary_csv}")
    print(f"{len(ranked_sas)} Solution Architect(s) scored. {impute_note}")


if __name__ == "__main__":
    main()
