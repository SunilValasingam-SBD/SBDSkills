#!/usr/bin/env python
"""
map_app_integrations.py

Given:
  1. An Excel workbook with an applications tab (default "ExcelApplications") listing
     applications, one per row.
  2. An integration catalog, which can come from EITHER:
       a. A JSON file (--catalog-json) containing the "values" array from the Ardoq
          "Integration Catalog" report (as returned by the local `ardoq` MCP server's
          get_integration_catalog / run_report tools), where each row is an object with
          keys such as "Initiator App ID", "Initiator App", "Terminator App ID",
          "Terminator", "Data Object", "Display Text", "Interfaces". "Initiator App" and
          "Terminator" may be lists of names.
       b. An Excel catalog tab (default "Integration Catalog", legacy mode) with an
          Initiator App column and a Terminator column.

Produce a new workbook with:
  - A "Summary" tab listing each application, its integration count, and its sheet name.
  - One tab per application containing every catalog row where that application is either
    the Initiator or the Terminator (matched by name and, when available, by App ID).

By default, applications with zero documented integrations are skipped entirely (no sheet,
no summary row). Use --include-empty to keep them (their sheet will contain only headers).

Usage (Ardoq MCP catalog, preferred):
    py map_app_integrations.py --input "C:\\path\\Workbook.xlsx" --catalog-json "C:\\path\\catalog.json"

Usage (legacy Excel catalog tab):
    py map_app_integrations.py --input "C:\\path\\Workbook.xlsx"

Optional arguments let you point at differently named sheets/columns and control output path.
"""

import argparse
import json
import os
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def sanitize_sheet_name(name, existing):
    """Return a unique, Excel-legal (<=31 chars, no []:*?/\\) sheet name."""
    clean = re.sub(r'[\[\]:*?/\\]', '-', str(name)).strip()
    if not clean:
        clean = "Sheet"
    base = clean[:31]
    candidate = base
    n = 1
    while candidate.lower() in existing:
        suffix = f"_{n}"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    existing.add(candidate.lower())
    return candidate


def find_column(headers, wanted, label):
    if wanted in headers:
        return headers.index(wanted)
    # case-insensitive fallback
    lower = [str(h).strip().lower() if h is not None else "" for h in headers]
    if wanted.strip().lower() in lower:
        return lower.index(wanted.strip().lower())
    raise SystemExit(
        f"Could not find column '{wanted}' for {label}. "
        f"Available columns: {headers}"
    )


def flatten_cell(value):
    """Render a cell value (possibly a list, from the Ardoq MCP JSON catalog) as a
    plain string suitable for an Excel cell."""
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value if v not in (None, ""))
    return value


def values_match(cell_value, target):
    """True if target (a string) matches cell_value, which may be a string or a
    list of strings (case-insensitive)."""
    if cell_value is None or target is None:
        return False
    target_norm = str(target).strip().lower()
    if not target_norm:
        return False
    if isinstance(cell_value, (list, tuple)):
        return any(str(v).strip().lower() == target_norm for v in cell_value)
    return str(cell_value).strip().lower() == target_norm


def load_catalog_json(path):
    """Load the Ardoq Integration Catalog from a JSON file.

    Accepts either the raw report object as returned by the ardoq MCP server's
    get_integration_catalog/run_report tools (``{"_meta": {...}, "values": [...]}``)
    or a bare list of row objects. Returns (headers, rows) where headers is an
    ordered list of column keys and rows is a list of dicts.
    """
    with open(path, encoding="utf-8") as f:
        data = json.load(f)

    if isinstance(data, dict) and "values" in data:
        rows = data["values"]
        meta_columns = ((data.get("_meta") or {}).get("columns")) or []
        headers = [c.get("key") or c.get("label") for c in meta_columns]
    elif isinstance(data, list):
        rows = data
        headers = []
    else:
        raise SystemExit(
            "Unrecognized --catalog-json format. Expected a report object with "
            "a 'values' list, or a bare list of row objects."
        )

    if not headers:
        # Fall back to the union of keys across rows, preserving first-seen order.
        seen = []
        for row in rows:
            for key in row.keys():
                if key not in seen:
                    seen.append(key)
        headers = seen

    return headers, rows


def main():
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "_shared"))
    from sa_paths import output_dir_for, latest_input

    parser = argparse.ArgumentParser(description="Map applications to their documented integrations.")
    parser.add_argument("--input", help="Path to the source Excel workbook (applications list). "
                         "Defaults to the newest .xlsx file in the shared SA\\Inputs folder.")
    parser.add_argument("--output", help="Path for the generated workbook. Defaults to "
                         "SA\\Outputs\\app-integration-mapper\\<input>_Integrations.xlsx.")
    parser.add_argument("--apps-sheet", default="ExcelApplications", help="Sheet name containing the application list.")
    parser.add_argument("--app-name-column", default="Application Name", help="Column in the applications sheet with the app name.")
    parser.add_argument("--app-id-column", default="App ID", help="Optional column in the applications sheet with the app's Ardoq App ID, used as an extra match key alongside the name. Ignored if not present.")

    parser.add_argument("--catalog-json", help="Path to a JSON file containing the Ardoq Integration Catalog "
                         "(the object returned by the ardoq MCP server's get_integration_catalog/run_report "
                         "tools, or a bare list of row objects). Preferred over --catalog-sheet.")
    parser.add_argument("--initiator-id-column", default="Initiator App ID", help="Catalog column/key with the initiator app's App ID (JSON catalog only).")
    parser.add_argument("--terminator-id-column", default="Terminator App ID", help="Catalog column/key with the terminator app's App ID (JSON catalog only).")

    parser.add_argument("--catalog-sheet", default="Integration Catalog", help="Legacy: sheet name containing the integration catalog, used only if --catalog-json is not given.")
    parser.add_argument("--initiator-column", default="Initiator App", help="Catalog column/key with the initiator app name.")
    parser.add_argument("--terminator-column", default="Terminator", help="Catalog column/key with the terminator app name.")

    parser.add_argument("--include-empty", action="store_true", help="Keep applications with zero integrations (empty sheet + summary row). Default is to skip them.")
    args = parser.parse_args()

    src = args.input or latest_input("*.xlsx")
    if not src:
        raise SystemExit("No --input given and no .xlsx file found in the shared SA\\Inputs folder.")
    src = str(src)
    if not os.path.isfile(src):
        raise SystemExit(f"Input file not found: {src}")

    if args.output:
        out = args.output
    else:
        base = os.path.splitext(os.path.basename(src))[0]
        out = str(output_dir_for("app-integration-mapper") / f"{base}_Integrations.xlsx")

    wb = openpyxl.load_workbook(src, data_only=True)
    if args.apps_sheet not in wb.sheetnames:
        raise SystemExit(f"Sheet '{args.apps_sheet}' not found. Available sheets: {wb.sheetnames}")

    apps_ws = wb[args.apps_sheet]
    apps_headers = [c.value for c in apps_ws[1]]
    app_name_idx = find_column(apps_headers, args.app_name_column, "the applications sheet")
    try:
        app_id_idx = find_column(apps_headers, args.app_id_column, "the applications sheet")
    except SystemExit:
        app_id_idx = None  # App ID column is optional

    apps = []  # list of (name, app_id_or_None)
    for row in apps_ws.iter_rows(min_row=2, values_only=True):
        name = row[app_name_idx]
        if name:
            app_id = row[app_id_idx] if app_id_idx is not None else None
            apps.append((str(name).strip(), str(app_id).strip() if app_id else None))

    # Load the catalog, either from Ardoq MCP JSON (preferred) or a legacy Excel sheet,
    # normalizing both into (cat_headers, cat_rows) where cat_rows is a list of dicts.
    if args.catalog_json:
        if not os.path.isfile(args.catalog_json):
            raise SystemExit(f"--catalog-json file not found: {args.catalog_json}")
        cat_headers, cat_rows = load_catalog_json(args.catalog_json)
        init_key = args.initiator_column if args.initiator_column in cat_headers else "Initiator App"
        term_key = args.terminator_column if args.terminator_column in cat_headers else "Terminator"
        init_id_key = args.initiator_id_column if args.initiator_id_column in cat_headers else None
        term_id_key = args.terminator_id_column if args.terminator_id_column in cat_headers else None
    else:
        if args.catalog_sheet not in wb.sheetnames:
            raise SystemExit(f"Sheet '{args.catalog_sheet}' not found. Available sheets: {wb.sheetnames}")
        cat_ws = wb[args.catalog_sheet]
        cat_headers = [c.value for c in cat_ws[1]]
        init_idx = find_column(cat_headers, args.initiator_column, "the catalog sheet")
        term_idx = find_column(cat_headers, args.terminator_column, "the catalog sheet")
        cat_rows = [dict(zip(cat_headers, r)) for r in cat_ws.iter_rows(min_row=2, values_only=True)]
        init_key, term_key = cat_headers[init_idx], cat_headers[term_idx]
        init_id_key = term_id_key = None

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    existing_names = set()
    summary = []

    for app, app_id in apps:
        matches = []
        for r in cat_rows:
            is_match = values_match(r.get(init_key), app) or values_match(r.get(term_key), app)
            if not is_match and app_id:
                if init_id_key and values_match(r.get(init_id_key), app_id):
                    is_match = True
                elif term_id_key and values_match(r.get(term_id_key), app_id):
                    is_match = True
            if is_match:
                matches.append(r)

        if not matches and not args.include_empty:
            continue

        sheet_name = sanitize_sheet_name(app, existing_names)
        ws = out_wb.create_sheet(sheet_name)
        ws.append(list(cat_headers))
        for c in ws[1]:
            c.font = copy(c.font)
            c.font = c.font.copy(bold=True)
        flat_rows = [[flatten_cell(r.get(h)) for h in cat_headers] for r in matches]
        for fr in flat_rows:
            ws.append(fr)
        for i, h in enumerate(cat_headers, start=1):
            col_letter = get_column_letter(i)
            max_len = max([len(str(h))] + [len(str(fr[i - 1])) if fr[i - 1] is not None else 0 for fr in flat_rows])
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)

        summary.append((app, len(matches), sheet_name))

    summary_ws = out_wb.create_sheet("Summary", 0)
    summary_ws.append(["Application Name", "Integration Count", "Sheet Name"])
    for c in summary_ws[1]:
        c.font = copy(c.font)
        c.font = c.font.copy(bold=True)
    for app, count, sheet in summary:
        summary_ws.append([app, count, sheet])
    summary_ws.column_dimensions['A'].width = 40
    summary_ws.column_dimensions['B'].width = 18
    summary_ws.column_dimensions['C'].width = 30

    out_wb.save(out)

    total_apps = len(apps)
    with_integrations = sum(1 for _, count, _ in summary if count > 0)
    print(f"Saved: {out}")
    print(f"Total applications in '{args.apps_sheet}': {total_apps}")
    print(f"Applications with documented integrations: {with_integrations}")
    print(f"Sheets created: {len(summary)}")


if __name__ == "__main__":
    main()
