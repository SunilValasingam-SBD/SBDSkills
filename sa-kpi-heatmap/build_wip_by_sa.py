"""Build a WIP Projects/Ideas by Solution Architect workbook.

Reads the raw work list CSV, filters to WIP Projects/Ideas (Work Ref starting
PRJ or ID), and writes an Excel workbook with:
  - "WIP Projects" tab: one row per project, sorted by SA then Start Date,
    with alternating fill color whenever the Solution Architect changes.
  - "Summary" tab: count of WIP projects/ideas per SA.
"""
import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Alternating band colors applied each time the SA changes.
BAND_COLORS = ["FFFFFF", "DCE6F1"]
HEADER_FILL = "4472C4"
HEADER_FONT_COLOR = "FFFFFF"

COLUMNS = [
    ("Solution Architect", "Work Ref: Solution Architect"),
    ("Work Ref", "Work Ref"),
    ("Project Name", "Work Ref: Name"),
    ("Program Owning Portfolio", "Work Ref: Program Owning Portfolio"),
    ("BRM", "Work Ref: BRM"),
    ("Start Date", "Work Ref: Start Date"),
    ("Finish Date", "Work Ref: Finish Date"),
]


def load_rows(input_path):
    with open(input_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
    return rows


def filter_wip(rows):
    out = []
    for r in rows:
        status = (r.get("Status") or "").strip()
        work_ref = (r.get("Work Ref") or "").strip()
        if status.upper() != "WIP":
            continue
        if not (work_ref.upper().startswith("PRJ") or work_ref.upper().startswith("ID")):
            continue
        out.append(r)
    return out


def sa_label(r):
    val = (r.get("Work Ref: Solution Architect") or "").strip()
    return val if val else "Unassigned"


def main():
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "_shared"))
    from sa_paths import output_dir_for, latest_input

    ap = argparse.ArgumentParser()
    ap.add_argument("--input", default=None, help="Path to the raw work list CSV. Defaults to the "
                    "newest .csv file in the shared SA\\Inputs folder.")
    ap.add_argument(
        "--output",
        default=None,
        help="Output Excel path (default: SA\\Outputs\\sa-kpi-heatmap\\wip_projects_by_sa.xlsx)",
    )
    args = ap.parse_args()

    input_path = Path(args.input) if args.input else latest_input("*.csv")
    if not input_path:
        raise SystemExit("No --input given and no .csv file found in the shared SA\\Inputs folder.")
    output_path = Path(args.output) if args.output else output_dir_for("sa-kpi-heatmap") / "wip_projects_by_sa.xlsx"

    rows = load_rows(input_path)
    wip_rows = filter_wip(rows)

    # Sort by SA (Unassigned last-ish is fine alphabetically), then Start Date string.
    def sort_key(r):
        return (sa_label(r), r.get("Work Ref: Start Date") or "")

    wip_rows.sort(key=sort_key)

    wb = Workbook()

    # --- WIP Projects tab ---
    ws = wb.active
    ws.title = "WIP Projects"

    headers = [c[0] for c in COLUMNS]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color=HEADER_FONT_COLOR)
        cell.fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    band_index = 0
    prev_sa = None
    sa_counts = {}
    for r in wip_rows:
        sa = sa_label(r)
        sa_counts[sa] = sa_counts.get(sa, 0) + 1
        if sa != prev_sa:
            band_index = (band_index + 1) % len(BAND_COLORS)
            prev_sa = sa
        fill_color = BAND_COLORS[band_index]
        values = [sa if key == "Work Ref: Solution Architect" else r.get(key, "") for _, key in COLUMNS]
        ws.append(values)
        row_idx = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_idx).fill = PatternFill(
                start_color=fill_color, end_color=fill_color, fill_type="solid"
            )

    # Column widths
    widths = [20, 10, 55, 35, 22, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # --- Summary tab ---
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Solution Architect", "WIP Projects/Ideas Count"])
    for col_idx in range(1, 3):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color=HEADER_FONT_COLOR)
        cell.fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for sa in sorted(sa_counts.keys()):
        ws2.append([sa, sa_counts[sa]])

    total_row = ws2.max_row + 1
    ws2.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    ws2.cell(row=total_row, column=2, value=sum(sa_counts.values())).font = Font(bold=True)

    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 24

    wb.save(output_path)
    print(f"Wrote {len(wip_rows)} WIP projects/ideas across {len(sa_counts)} SAs to {output_path}")


if __name__ == "__main__":
    main()
